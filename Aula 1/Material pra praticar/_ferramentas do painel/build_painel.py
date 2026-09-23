# -*- coding: utf-8 -*-
"""
Monta o arquivo HTML final e autocontido do painel da Aurora a partir de:
  - template.html          (nesta mesma pasta; já deve ter o texto da aba
                             Insights atualizado à mão antes de rodar isto)
  - chart.umd.min.js       (nesta mesma pasta)
  - chartjs-plugin-datalabels.min.js (nesta mesma pasta)
  - fonts_inline.css       (nesta mesma pasta)
  - a planilha Controle_Recebimentos.xlsx (caminho passado por parâmetro)

Lê a aba "Recebimentos" e a aba "Cadastro de Clientes" diretamente da
planilha (não precisa de nenhum rows.json separado), monta o JSON de dados
embutido e gera o HTML final, 100% offline (sem nenhum <script src=...> ou
<link> externo).

Uso:
    python3 build_painel.py "/caminho/Controle_Recebimentos.xlsx" "/caminho/saida/Aurora - Painel de Recebimentos.html"
"""
import json
import sys
import os
import openpyxl


def load_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb['Recebimentos']
    cadastro_ws = wb['Cadastro de Clientes']

    cli_map = {}
    for row in cadastro_ws.iter_rows(min_row=5, values_only=True):
        if not row or not row[0]:
            continue
        cli_map[row[0]] = (row[1], row[2])

    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == 'Data':
            header_row = i
            break
    if header_row is None:
        raise SystemExit('Não encontrei a linha de cabeçalho (Data | Cliente | ...) na aba Recebimentos.')

    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or row[0] is None:
            continue
        data, cliente, valor, forma = row[0], row[1], row[2], row[3]
        regiao, categoria = cli_map.get(cliente, (None, None))
        rows.append({
            'd': data.isoformat(),
            'm': data.month,
            'c': cliente,
            'v': float(valor),
            'p': forma or '',
            'r': regiao or 'Não identificado',
            'g': categoria or 'Não identificado',
        })
    return rows


def main():
    if len(sys.argv) != 3:
        raise SystemExit('Uso: python3 build_painel.py "/caminho/Controle_Recebimentos.xlsx" "/caminho/saida.html"')
    xlsx_path, out_path = sys.argv[1], sys.argv[2]
    here = os.path.dirname(os.path.abspath(__file__))

    rows = load_rows(xlsx_path)
    rows_json = json.dumps(rows, ensure_ascii=False)

    tpl = open(os.path.join(here, 'template.html'), encoding='utf-8').read()
    font_css = open(os.path.join(here, 'fonts_inline.css'), encoding='utf-8').read()
    chartjs = open(os.path.join(here, 'chart.umd.min.js'), encoding='utf-8').read()
    datalabels = open(os.path.join(here, 'chartjs-plugin-datalabels.min.js'), encoding='utf-8').read()

    out = tpl.replace('__ROWS_JSON__', rows_json)

    old_fonts = '''<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;500;600;700&family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">'''
    if old_fonts in tpl:
        out = out.replace(old_fonts, f'<style>\n{font_css}\n</style>')

    old_scripts = '''<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.4/chart.umd.min.js"></script>
<script src="https://cdnjs.cloudflare.com/ajax/libs/chartjs-plugin-datalabels/2.2.0/chartjs-plugin-datalabels.min.js"></script>'''
    if old_scripts in tpl:
        out = out.replace(old_scripts, f'<script>\n{chartjs}\n</script>\n<script>\n{datalabels}\n</script>')

    if '<script src=' in out or 'fonts.googleapis' in out or 'cdnjs.cloudflare' in out:
        raise SystemExit('ATENÇÃO: sobrou referência externa no HTML final — não deve ser entregue assim (o arquivo tem que funcionar 100% offline).')

    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(out)
    print(f'OK: {out_path} ({len(out)} bytes, {len(rows)} recebimentos)')


if __name__ == '__main__':
    main()
